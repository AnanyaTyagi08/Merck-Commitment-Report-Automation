"""Excel report generator for Merck Commitment Reports."""

from copy import copy
from pathlib import Path
from typing import Union

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def apply_sheet_layout(
    ws: Worksheet,
    header_row: int = 1,
) -> None:
    """Freeze rows through the header and enable column filters."""
    if header_row < 1 or header_row > ws.max_row:
        raise ValueError("Header row must exist in the worksheet.")

    ws.freeze_panes = f"A{header_row + 1}"

    last_column = get_column_letter(ws.max_column)
    ws.auto_filter.ref = (
        f"A{header_row}:{last_column}{ws.max_row}"
    )

def apply_font_color(
    ws: Worksheet,
    cell_addresses: list[str],
    color: str,
) -> None:
    """Color specified cells red or purple, preserving other font settings."""
    colors = {
        "red": "FFFF0000",
        "purple": "FF800080",
    }

    if color not in colors:
        raise ValueError("Color must be 'red' or 'purple'.")

    for address in cell_addresses:
        cell = ws[address]
        updated_font = copy(cell.font)
        updated_font.color = colors[color]
        cell.font = updated_font

def build_grt_spmt_workbook(
    df: pd.DataFrame,
    output_path: Union[str, Path],
) -> None:
    """Generate the formatted Excel workbook with required tabs and views."""
    raise NotImplementedError("Pending implementation")

#implementing red overdue row labeling function for Task-12
def highlight_overdue_rows(ws: Worksheet, df: pd.DataFrame, header_row: int = 1) -> None:
    """Apply red font to all cells in rows where the Working Due Date is overdue."""
    today = pd.Timestamp.today().normalize()
    red_cells = []

    #enumeration starts at header_row + 1 because Excel is 1-indexed and has a header
    for row_idx, (_, row) in enumerate(df.iterrows(), start=header_row + 1):
        date_val = pd.to_datetime(row.get("Working Due Date"), errors="coerce")

        if pd.notna(date_val) and date_val < today:
            #grabbing every column's letter for this specific row index
            for col_idx in range(1, len(df.columns) + 1):
                col_letter = get_column_letter(col_idx)
                red_cells.append(f"{col_letter}{row_idx}")

    if red_cells:
        apply_font_color(ws, red_cells, "red")
