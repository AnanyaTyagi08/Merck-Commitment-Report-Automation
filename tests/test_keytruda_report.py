import pandas as pd
from openpyxl import load_workbook

from src.excel_generator import build_keytruda_workbook
from src.transforms import filter_keytruda_report


def test_keytruda_filter_rules():
    df = pd.DataFrame(
        [
            ("overdue", "pembrolizumab", "Ongoing", "", "Canada", "2026-07-31"),
            ("boundary", "pembrolizumab", "Ongoing", "", "Canada", "2027-02-28"),
            ("outside", "pembrolizumab", "Ongoing", "", "Canada", "2027-03-01"),
            ("placeholder98", "pembrolizumab", "Ongoing", "", "Canada", "2098-01-01"),
            ("placeholder99", "pembrolizumab", "Ongoing", "", "Canada", "2099-12-31"),
            ("completed", "pembrolizumab", "Completed", "", "Canada", "2026-10-01"),
            ("non_us_submitted", "pembrolizumab", "Submitted", "Submitted",
             "Canada", "2026-10-01"),
            ("us_short", "pembrolizumab", "Submitted", "Submitted",
             "United States", "2026-10-01"),
            ("us_long", "pembrolizumab", "Submitted", "Submitted",
             "United States of America", "2026-10-01"),
            ("other_product", "other", "Ongoing", "", "Canada", "2026-10-01"),
            ("missing_date", "pembrolizumab", "Ongoing", "", "Canada", None),
        ],
        columns=[
            "Commitment ID",
            "Primary Product Family",
            "Commitment Lifecycle State",
            "Completed State Qualifier",
            "Country Name",
            "Working Due Date",
        ],
    )
    df["Commitment Due Date"] = df["Working Due Date"]
    original = df.copy(deep=True)

    result = filter_keytruda_report(df, pd.Timestamp("2026-09-01"))

    assert result["Commitment ID"].tolist() == [
        "overdue",
        "us_short",
        "us_long",
        "boundary",
        "placeholder98",
        "placeholder99",
    ]
    pd.testing.assert_frame_equal(df, original)


def test_keytruda_due_date_fallback():
    df = pd.DataFrame(
        {
            "Primary Product Family": [" Pembrolizumab "],
            "Commitment Lifecycle State": ["Ongoing"],
            "Completed State Qualifier": [""],
            "Country Name": ["Canada"],
            "Working Due Date": [None],
            "Commitment Due Date": [pd.Timestamp("2026-12-31")],
        }
    )

    result = filter_keytruda_report(df, pd.Timestamp("2026-09-01"))

    assert len(result) == 1
    assert result.iloc[0]["Working Due Date"] == pd.Timestamp("2026-12-31")


def test_keytruda_excel_export(tmp_path):
    df = pd.DataFrame(
        {
            "Commitment ID": ["TEST-001"],
            "Working Due Date": [pd.Timestamp("2026-12-31")],
        }
    )

    output = build_keytruda_workbook(
        df, tmp_path, pd.Timestamp("2026-09-01")
    )

    assert output.name == "Commitments Report-Keytruda-Ongoing-Sep2026.xlsx"

    wb = load_workbook(output)
    try:
        ws = wb["Keytruda Sep 2026"]
        assert ws.freeze_panes == "A3"
        assert ws.auto_filter.ref == "A2:B3"
        assert ws["A2"].value == "Commitment ID"
        assert ws["A3"].value == "TEST-001"
        assert pd.Timestamp(ws["B3"].value) == pd.Timestamp("2026-12-31")
    finally:
        wb.close()
