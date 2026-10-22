import openpyxl
import pandas as pd
import pytest

from src.excel_generator import apply_font_color, apply_sheet_layout, build_grt_spmt_workbook


#skipped purposefully: out of scope for Task-6. excel workbook formatting
#(Red-Late/Purple-Pending tabs) will be handled in a separate epic/ticket
@pytest.mark.skip(reason="Pending excel generator implementation")
def test_required_workbook_tabs_and_freeze_panes(tmp_path):
    """Verify workbook generates Red-Late and Purple-Pending tabs with frozen headers."""
    sample_df = pd.DataFrame({
        "Commitment Lifecycle State": ["Ongoing"],
        "Metric Indicator": ["Purple"],
        "Working Due Date": [pd.Timestamp.now() - pd.Timedelta(days=10)],
        "Country Name": ["United States"]
    })

    output_path = tmp_path / "test_report.xlsx"
    build_grt_spmt_workbook(sample_df, output_path)

    wb = openpyxl.load_workbook(output_path)
    sheet_names = wb.sheetnames

    #asserting expected sheets exist per RPP-200-ER03
    assert "Red-Late" in sheet_names
    assert "Purple-Pending" in sheet_names

    #verifying top row is frozen on main sheet
    ws = wb.active
    assert ws.freeze_panes == "A2"
@pytest.mark.parametrize("header_row", [1, 2])
def test_sheet_layout_survives_save(tmp_path, header_row):
    """Verify frozen headers and filters persist in the saved workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active

    if header_row == 2:
        ws.append(["Monthly report"])

    ws.append(["Commitment ID", "Metric Indicator"])
    ws.append(["TEST-001", "Green"])

    apply_sheet_layout(ws, header_row=header_row)

    output_path = tmp_path / "layout_test.xlsx"
    wb.save(output_path)
    wb.close()

    saved_wb = openpyxl.load_workbook(output_path)
    saved_ws = saved_wb.active

    assert saved_ws.freeze_panes == f"A{header_row + 1}"
    assert saved_ws.auto_filter.ref == (
        f"A{header_row}:B{header_row + 1}"
    )
    assert saved_ws.cell(header_row + 1, 1).value == "TEST-001"

@pytest.mark.parametrize(
    "color, expected",
    [
        ("red", "FFFF0000"),
        ("purple", "FF800080"),
    ],
)
def test_font_color_survives_save(tmp_path, color, expected):
    """Verify color persists without changing other font settings or cells."""
    wb = openpyxl.Workbook()
    ws = wb.active

    ws["A1"] = "Target"
    ws["A1"].font = openpyxl.styles.Font(
        name="Calibri",
        size=12,
        bold=True,
        italic=True,
    )
    ws["B1"] = "Unchanged"
    ws["B1"].font = openpyxl.styles.Font(color="FF000000")

    apply_font_color(ws, ["A1"], color)

    output_path = tmp_path / "font_test.xlsx"
    wb.save(output_path)
    wb.close()

    saved_wb = openpyxl.load_workbook(output_path)
    saved_ws = saved_wb.active
    font = saved_ws["A1"].font

    assert font.color.rgb == expected
    assert font.name == "Calibri"
    assert font.sz == 12
    assert font.bold is True
    assert font.italic is True
    assert saved_ws["A1"].value == "Target"
    assert saved_ws["B1"].font.color.rgb == "FF000000"

    saved_wb.close()
